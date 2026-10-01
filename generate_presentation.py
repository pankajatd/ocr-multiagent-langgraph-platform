"""
PowerPoint Presentation Generator for LangGraph Multi-Agent OCR System.
Creates a 15-slide corporate-grade 16:9 widescreen deck.

Generates: OCR_MultiAgent_LangGraph_Platform.pptx

Design principles:
  • Story-driven structure (Problem → Solution → How It Works → Results)
  • One idea per slide
  • Visual agent flow with numbered steps
  • Before/After examples for self-healing
  • Minimal text, maximum clarity
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ──────────────────── COLOR PALETTE ────────────────────
NAVY        = RGBColor(15, 23, 42)       # Slate 900
DARK_BLUE   = RGBColor(30, 58, 138)      # Blue 800
MID_BLUE    = RGBColor(37, 99, 235)      # Blue 600
LIGHT_BLUE  = RGBColor(59, 130, 246)     # Blue 500
SKY         = RGBColor(186, 230, 253)    # Sky 200
CYAN        = RGBColor(2, 132, 199)      # Cyan 600
EMERALD     = RGBColor(5, 150, 105)      # Emerald 600
AMBER       = RGBColor(217, 119, 6)      # Amber 600
ROSE        = RGBColor(225, 29, 72)      # Rose 600
PURPLE      = RGBColor(124, 58, 237)     # Violet 600
WHITE       = RGBColor(255, 255, 255)
OFF_WHITE   = RGBColor(248, 250, 252)    # Slate 50
CARD_BG     = RGBColor(255, 255, 255)
LIGHT_GRAY  = RGBColor(241, 245, 249)    # Slate 100
BORDER      = RGBColor(203, 213, 225)    # Slate 300
TEXT_DARK   = RGBColor(15, 23, 42)       # Slate 900
TEXT_BODY   = RGBColor(51, 65, 85)       # Slate 700
TEXT_MUTED  = RGBColor(100, 116, 139)    # Slate 500
TEXT_LIGHT  = RGBColor(226, 232, 240)    # Slate 200

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def create_deck(output_filename="OCR_MultiAgent_LangGraph_Platform.pptx"):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ── Helper: full-slide background ──
    def fill_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        bg.rotation = 0.0

    # ── Helper: accent bar at top of slide ──
    def top_bar(slide, color=MID_BLUE, height=Inches(0.08)):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

    # ── Helper: footer stripe ──
    def footer(slide, text="LangGraph Multi-Agent OCR Platform  |  github.com/pankajatd/ocr-multiagent-langgraph-platform"):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SLIDE_H - Inches(0.45), SLIDE_W, Inches(0.45))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(0.6), SLIDE_H - Inches(0.40), Inches(12), Inches(0.35))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = text
        p.font.size = Pt(9); p.font.color.rgb = TEXT_LIGHT; p.font.italic = True

    # ── Helper: section header on content slides ──
    def section_header(slide, category, title, subtitle=None):
        top_bar(slide)
        # Category tag
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11), Inches(0.35))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = category.upper()
        p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = MID_BLUE
        p.font.name = "Calibri"
        # Title
        tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11), Inches(0.7))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = title
        p2.font.size = Pt(26); p2.font.bold = True; p2.font.color.rgb = NAVY
        p2.font.name = "Calibri"
        # Subtitle
        if subtitle:
            tb3 = slide.shapes.add_textbox(Inches(0.8), Inches(1.30), Inches(11), Inches(0.5))
            tf3 = tb3.text_frame; tf3.word_wrap = True
            p3 = tf3.paragraphs[0]; p3.text = subtitle
            p3.font.size = Pt(14); p3.font.color.rgb = TEXT_MUTED
            p3.font.name = "Calibri"

    # ── Helper: rounded card with icon header ──
    def card(slide, left, top, width, height, icon, title, bullets, accent=MID_BLUE, bullet_size=Pt(12)):
        # Card background
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid(); shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = BORDER; shape.line.width = Pt(1)
        # Accent top bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.07))
        bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
        # Icon + Title
        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.18), width - Inches(0.5), Inches(0.45))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = f"{icon}  {title}"
        p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = NAVY
        p.font.name = "Calibri"
        # Bullet content
        tb2 = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.65), width - Inches(0.5), height - Inches(0.85))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        for i, item in enumerate(bullets):
            p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
            p.text = f"▸  {item}"
            p.font.size = bullet_size; p.font.color.rgb = TEXT_BODY
            p.font.name = "Calibri"; p.space_after = Pt(5)

    # ── Helper: numbered step box (for flow diagrams) ──
    def step_box(slide, left, top, width, height, number, label, color=MID_BLUE, sublabel=None):
        # Circle with number
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.55), Inches(0.55))
        circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.line.fill.background()
        tf = circle.text_frame; tf.word_wrap = False
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        p = tf.paragraphs[0]; p.text = str(number)
        p.font.size = Pt(16); p.font.bold = True; p.font.color.rgb = WHITE
        p.font.name = "Calibri"
        # Label box below
        tb = slide.shapes.add_textbox(left - Inches(0.15), top + Inches(0.65), width + Inches(0.3), height)
        tf2 = tb.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = label; p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(11); p2.font.bold = True; p2.font.color.rgb = NAVY
        p2.font.name = "Calibri"
        if sublabel:
            p3 = tf2.add_paragraph(); p3.text = sublabel; p3.alignment = PP_ALIGN.CENTER
            p3.font.size = Pt(9); p3.font.color.rgb = TEXT_MUTED
            p3.font.name = "Calibri"

    # ── Helper: arrow connector (simple right arrow) ──
    def arrow_right(slide, left, top, length=Inches(0.6), color=BORDER):
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, length, Inches(0.22))
        arr.fill.solid(); arr.fill.fore_color.rgb = color; arr.line.fill.background()

    # ── Helper: arrow connector (down arrow) ──
    def arrow_down(slide, left, top, length=Inches(0.5), color=BORDER):
        arr = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, Inches(0.22), length)
        arr.fill.solid(); arr.fill.fore_color.rgb = color; arr.line.fill.background()

    # ── Helper: big stat number ──
    def stat_block(slide, left, top, number_text, label_text, color=MID_BLUE):
        tb = slide.shapes.add_textbox(left, top, Inches(2.5), Inches(0.6))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = number_text
        p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = color
        p.font.name = "Calibri"
        tb2 = slide.shapes.add_textbox(left, top + Inches(0.55), Inches(2.5), Inches(0.4))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = label_text
        p2.font.size = Pt(12); p2.font.color.rgb = TEXT_MUTED
        p2.font.name = "Calibri"

    # ═══════════════════════════════════════════════════════
    # SLIDE 1: TITLE SLIDE
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    fill_bg(s, NAVY)

    # Decorative diagonal accent
    diag = s.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(8.5), 0, Inches(5.5), SLIDE_H)
    diag.fill.solid(); diag.fill.fore_color.rgb = DARK_BLUE; diag.line.fill.background()

    # Main title
    tb = s.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(7.5), Inches(1.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "MULTI-AGENT OCR PLATFORM"
    p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = SKY
    p.font.name = "Calibri"; p.space_after = Pt(6)

    tb2 = s.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(7.5), Inches(2.0))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p2 = tf2.paragraphs[0]; p2.text = "Intelligent Document\nProcessing with\nSelf-Healing AI Agents"
    p2.font.size = Pt(42); p2.font.bold = True; p2.font.color.rgb = WHITE
    p2.font.name = "Calibri"; p2.line_spacing = Pt(50)

    tb3 = s.shapes.add_textbox(Inches(1.0), Inches(4.6), Inches(7.5), Inches(0.8))
    tf3 = tb3.text_frame; tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "Built with LangGraph  •  RapidOCR  •  Streamlit"
    p3.font.size = Pt(16); p3.font.color.rgb = TEXT_LIGHT
    p3.font.name = "Calibri"

    # Right side: 3 domain icons
    for i, (icon, label) in enumerate([("📄", "Documents"), ("🚗", "License Plates"), ("🧾", "Invoices")]):
        y = Inches(2.0) + Inches(1.5) * i
        tb = s.shapes.add_textbox(Inches(10.0), y, Inches(2.5), Inches(1.0))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = f"{icon}  {label}"
        p.font.size = Pt(18); p.font.color.rgb = WHITE; p.font.name = "Calibri"

    # Author line
    tb4 = s.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(7.5), Inches(0.5))
    tf4 = tb4.text_frame; tf4.word_wrap = True
    p4 = tf4.paragraphs[0]; p4.text = "Pankaj Todkar  |  github.com/pankajatd"
    p4.font.size = Pt(12); p4.font.color.rgb = TEXT_MUTED; p4.font.name = "Calibri"

    # ═══════════════════════════════════════════════════════
    # SLIDE 2: AGENDA / TABLE OF CONTENTS
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Overview", "What We'll Cover Today")
    footer(s)

    agenda_items = [
        ("01", "The Problem", "Why traditional OCR fails on real-world documents"),
        ("02", "Our Solution", "Multi-agent architecture with intelligent routing"),
        ("03", "Meet the Agents", "8 specialized AI agents and what each one does"),
        ("04", "The Complete Flow", "Step-by-step: from image upload to structured output"),
        ("05", "Self-Healing Magic", "How agents automatically fix their own mistakes"),
        ("06", "Three Business Domains", "Documents, License Plates, and Invoices"),
        ("07", "Live Demo & Results", "Dashboard, CLI, and test verification"),
    ]
    for i, (num, title, desc) in enumerate(agenda_items):
        y = Inches(1.85) + Inches(0.72) * i
        # Number circle
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), y, Inches(0.48), Inches(0.48))
        circ.fill.solid(); circ.fill.fore_color.rgb = MID_BLUE; circ.line.fill.background()
        ctf = circ.text_frame; ctf.paragraphs[0].alignment = PP_ALIGN.CENTER
        ctf.paragraphs[0].text = num; ctf.paragraphs[0].font.size = Pt(13)
        ctf.paragraphs[0].font.bold = True; ctf.paragraphs[0].font.color.rgb = WHITE
        # Title
        tb = s.shapes.add_textbox(Inches(1.7), y, Inches(3), Inches(0.4))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title
        p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = NAVY; p.font.name = "Calibri"
        # Description
        tb2 = s.shapes.add_textbox(Inches(4.8), y + Inches(0.02), Inches(7), Inches(0.4))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = desc
        p2.font.size = Pt(13); p2.font.color.rgb = TEXT_MUTED; p2.font.name = "Calibri"

    # ═══════════════════════════════════════════════════════
    # SLIDE 3: THE PROBLEM
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "The Challenge", "Why Traditional OCR Fails",
                   "Single-pass OCR tools crash or produce garbage on real-world messy documents.")
    footer(s)

    problems = [
        ("❌", "One-Size-Fits-All", "Treats a book page, license plate, and invoice identically — missing domain-specific rules.",
         ROSE),
        ("❌", "No Error Recovery", "A single misread character ($275 → $2750) corrupts the entire output with no fix.",
         ROSE),
        ("❌", "Manual Intervention", "Humans must review and correct every OCR error — expensive and slow.",
         ROSE),
        ("❌", "Fragile Pipeline", "If preprocessing fails (dark image, skewed scan), the whole pipeline crashes.",
         ROSE),
    ]
    for i, (icon, title, desc, color) in enumerate(problems):
        y = Inches(2.1) + Inches(1.2) * i
        # Icon box
        ib = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y, Inches(0.55), Inches(0.55))
        ib.fill.solid(); ib.fill.fore_color.rgb = RGBColor(254, 226, 226); ib.line.fill.background()
        itf = ib.text_frame; itf.paragraphs[0].alignment = PP_ALIGN.CENTER
        itf.paragraphs[0].text = icon; itf.paragraphs[0].font.size = Pt(16)
        # Title + Description
        tb = s.shapes.add_textbox(Inches(1.8), y - Inches(0.02), Inches(10), Inches(0.3))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title
        p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = NAVY; p.font.name = "Calibri"
        tb2 = s.shapes.add_textbox(Inches(1.8), y + Inches(0.30), Inches(10), Inches(0.5))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = desc
        p2.font.size = Pt(13); p2.font.color.rgb = TEXT_BODY; p2.font.name = "Calibri"

    # ═══════════════════════════════════════════════════════
    # SLIDE 4: OUR SOLUTION (HIGH LEVEL)
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "The Solution", "A Team of AI Agents That Work Together",
                   "Instead of one monolithic OCR tool, we use 8 specialized agents that collaborate, route intelligently, and self-correct.")
    footer(s)

    solutions = [
        ("🧠", "Smart Routing", "An Orchestrator Agent reads the image and decides: Is this a document, a license plate, or an invoice? Then sends it to the right specialist.", MID_BLUE),
        ("🎯", "Domain Experts", "Each specialist agent knows its domain deeply — accounting rules for invoices, traffic syntax for plates, layout logic for documents.", EMERALD),
        ("🔄", "Self-Healing", "When an agent detects an error (wrong math, confused characters), an Error Resolver Agent automatically fixes it — no human needed.", AMBER),
        ("🛡️", "Never Crashes", "Built-in retry limits and graceful degradation ensure the system always returns a result, even on terrible image quality.", PURPLE),
    ]
    for i, (icon, title, desc, color) in enumerate(solutions):
        col = i % 2
        row = i // 2
        x = Inches(0.8) + Inches(6.0) * col
        y = Inches(2.1) + Inches(2.4) * row
        card(s, x, y, Inches(5.6), Inches(2.1), icon, title, [desc], accent=color)

    # ═══════════════════════════════════════════════════════
    # SLIDE 5: MEET THE 8 AGENTS (ROSTER)
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Agent Roster", "Meet the 8 AI Agents",
                   "Each agent has a single responsibility and passes state to the next agent via LangGraph.")
    footer(s)

    agents = [
        ("🛠️", "Preprocessor", "Cleans & enhances the image\n(deskew, contrast, noise removal)", MID_BLUE),
        ("👁️", "OCR Engine", "Reads text from the image\n(RapidOCR + Tesseract fallback)", CYAN),
        ("🧠", "Orchestrator", "Classifies document type &\nroutes to the right specialist", DARK_BLUE),
        ("📄", "Doc Digitizer", "Extracts layout, paragraphs,\nheadings & search index", MID_BLUE),
        ("🚗", "Plate Reader", "Extracts license plate number,\nvalidates jurisdiction format", EMERALD),
        ("🧾", "Invoice Scanner", "Extracts vendor, line items,\ntotals & verifies math", AMBER),
        ("🩺", "Error Resolver", "Diagnoses & auto-fixes errors\n(math, char confusion, re-filter)", ROSE),
        ("📦", "Postprocessor", "Packages final output as\nstructured JSON & Markdown", PURPLE),
    ]
    for i, (icon, name, desc, color) in enumerate(agents):
        col = i % 4
        row = i // 4
        x = Inches(0.6) + Inches(3.1) * col
        y = Inches(2.0) + Inches(2.6) * row
        # Card
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.8), Inches(2.2))
        shape.fill.solid(); shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = BORDER; shape.line.width = Pt(1)
        # Accent bar
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(2.8), Inches(0.06))
        bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
        # Icon
        tb = s.shapes.add_textbox(x + Inches(0.15), y + Inches(0.2), Inches(2.5), Inches(0.5))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = f"{icon}  {name}"
        p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = NAVY; p.font.name = "Calibri"
        # Description
        tb2 = s.shapes.add_textbox(x + Inches(0.15), y + Inches(0.8), Inches(2.5), Inches(1.2))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = desc
        p2.font.size = Pt(11); p2.font.color.rgb = TEXT_BODY; p2.font.name = "Calibri"
        p2.line_spacing = Pt(16)

    # ═══════════════════════════════════════════════════════
    # SLIDE 6: THE COMPLETE PIPELINE FLOW (VISUAL DIAGRAM)
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "How It Works", "The Complete Processing Pipeline",
                   "Follow the numbered steps to see how an image travels through all 8 agents.")
    footer(s)

    # Row 1: Steps 1-4
    positions_row1 = [
        (Inches(0.7), "Upload", "Image/PDF/\nCSV/TXT", MID_BLUE),
        (Inches(3.2), "Preprocess", "Deskew, CLAHE\nNoise cleanup", MID_BLUE),
        (Inches(5.7), "OCR Scan", "Extract text &\nconfidence scores", CYAN),
        (Inches(8.2), "Classify", "Detect document\ntype & route", DARK_BLUE),
    ]
    for i, (x, label, sub, color) in enumerate(positions_row1, 1):
        step_box(s, x, Inches(2.0), Inches(1.8), Inches(0.8), i, label, color, sub)
        if i < 4:
            arrow_right(s, x + Inches(1.7), Inches(2.18), Inches(0.5), BORDER)

    # Branch arrows going down from step 4
    arrow_down(s, Inches(8.42), Inches(3.6), Inches(0.5), BORDER)

    # Row 2: Steps 5a, 5b, 5c (specialist branches)
    specialists = [
        (Inches(1.5), "5a", "Document\nDigitizer", "Layout &\nsearch index", MID_BLUE),
        (Inches(5.2), "5b", "License Plate\nReader", "ANPR &\nvalidation", EMERALD),
        (Inches(8.9), "5c", "Invoice\nScanner", "Line items &\nmath audit", AMBER),
    ]

    # Horizontal line from step 4 center to each specialist
    branch_bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.72), Inches(4.1), Inches(7.9), Inches(0.04))
    branch_bar.fill.solid(); branch_bar.fill.fore_color.rgb = BORDER; branch_bar.line.fill.background()

    for x, num, label, sub, color in specialists:
        arrow_down(s, x + Inches(0.16), Inches(4.1), Inches(0.4), BORDER)
        step_box(s, x, Inches(4.5), Inches(1.8), Inches(0.8), num, label, color, sub)

    # Row 3: Steps 6 and 7
    step_box(s, Inches(3.5), Inches(6.0), Inches(1.8), Inches(0.5), 6, "Validate", ROSE, "Check accuracy\n& math")
    arrow_right(s, Inches(5.4), Inches(6.18), Inches(0.5), BORDER)
    step_box(s, Inches(6.3), Inches(6.0), Inches(1.8), Inches(0.5), 7, "Package", PURPLE, "JSON &\nMarkdown output")

    # Self-healing loop indicator
    loop_label = s.shapes.add_textbox(Inches(9.5), Inches(6.0), Inches(3.5), Inches(0.9))
    ltf = loop_label.text_frame; ltf.word_wrap = True
    lp = ltf.paragraphs[0]; lp.text = "🔄 If errors found →"
    lp.font.size = Pt(11); lp.font.bold = True; lp.font.color.rgb = ROSE; lp.font.name = "Calibri"
    lp2 = ltf.add_paragraph(); lp2.text = "Error Resolver auto-fixes\nor loops back to Step 2"
    lp2.font.size = Pt(10); lp2.font.color.rgb = TEXT_MUTED; lp2.font.name = "Calibri"

    # ═══════════════════════════════════════════════════════
    # SLIDE 7: ORCHESTRATOR AGENT DEEP DIVE
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Agent Deep Dive", "🧠 The Orchestrator Agent — Smart Document Routing",
                   "The Orchestrator analyzes the OCR text and automatically decides which specialist should handle it.")
    footer(s)

    # Three routing paths
    routes = [
        ("📄", "Routes to Document Digitizer", [
            "Detects keywords: paragraph, section, chapter, abstract",
            "High word count (>50 words)",
            "Normal aspect ratio (standard page)",
            "Example: Scanned research paper, letter, report",
        ], MID_BLUE),
        ("🚗", "Routes to License Plate Agent", [
            "Detects keywords: plate, registration, vehicle",
            "Very low word count (< 10 words)",
            "Wide aspect ratio (typical plate shape)",
            "Example: Highway toll camera capture",
        ], EMERALD),
        ("🧾", "Routes to Invoice Scanner", [
            "Detects keywords: invoice, total, subtotal, tax, amount, qty",
            "Medium word count with financial terms",
            "Contains number patterns like $, %, decimals",
            "Example: Vendor bill, bank statement, receipt",
        ], AMBER),
    ]
    for i, (icon, title, bullets, color) in enumerate(routes):
        x = Inches(0.6) + Inches(4.1) * i
        card(s, x, Inches(2.1), Inches(3.8), Inches(4.3), icon, title, bullets, accent=color)

    # ═══════════════════════════════════════════════════════
    # SLIDE 8: DOMAIN 1 — DOCUMENT DIGITIZER
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Business Domain 1", "📄 Document Digitizer Agent",
                   "Converts scanned papers into structured, searchable digital archives.")
    footer(s)

    card(s, Inches(0.6), Inches(2.1), Inches(5.8), Inches(4.5), "📝", "What It Extracts", [
        "Document title and section headings",
        "Body paragraphs in correct reading order",
        "Word count and estimated reading time",
        "Full-text search index (keyword → frequency map)",
        "Clean Markdown export for web publishing",
    ], accent=MID_BLUE)

    card(s, Inches(6.8), Inches(2.1), Inches(5.8), Inches(4.5), "💡", "Real-World Use Cases", [
        "Digitize paper archives for law firms & hospitals",
        "Convert scanned research papers to searchable text",
        "Index legacy documents for enterprise search engines",
        "Enable text-based search across thousands of PDFs",
        "Automated compliance document processing",
    ], accent=CYAN)

    # ═══════════════════════════════════════════════════════
    # SLIDE 9: DOMAIN 2 — LICENSE PLATE
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Business Domain 2", "🚗 License Plate Reader (ANPR)",
                   "Automatic Number Plate Recognition for smart traffic and security systems.")
    footer(s)

    card(s, Inches(0.6), Inches(2.1), Inches(5.8), Inches(4.5), "🔍", "What It Extracts", [
        "Alphanumeric plate registration number",
        "Jurisdiction validation (US, UK, EU formats)",
        "Character disambiguation (O→0, I→1, B→8)",
        "Filters out frame noise and dealer logos",
        "Simulated traffic telemetry (toll gate, lane, speed)",
    ], accent=EMERALD)

    card(s, Inches(6.8), Inches(2.1), Inches(5.8), Inches(4.5), "💡", "Real-World Use Cases", [
        "Highway toll collection (automatic billing)",
        "Parking lot entry/exit management",
        "Law enforcement vehicle lookups",
        "Smart city traffic flow monitoring",
        "Security checkpoint vehicle screening",
    ], accent=MID_BLUE)

    # ═══════════════════════════════════════════════════════
    # SLIDE 10: DOMAIN 3 — INVOICE SCANNER
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Business Domain 3", "🧾 Invoice & Receipt Scanner",
                   "Automated financial data extraction with mathematical integrity verification.")
    footer(s)

    card(s, Inches(0.6), Inches(2.1), Inches(5.8), Inches(4.5), "🔍", "What It Extracts", [
        "Vendor name and invoice ID",
        "Invoice date and payment terms",
        "Line items table (description, qty, price, total)",
        "Financial summary: subtotal, tax, grand total",
        "Math audit: verifies Subtotal + Tax = Grand Total",
    ], accent=AMBER)

    card(s, Inches(6.8), Inches(2.1), Inches(5.8), Inches(4.5), "💡", "Real-World Use Cases", [
        "Automated accounts payable data entry",
        "Expense report processing and validation",
        "Fraud detection via math discrepancy alerts",
        "ERP system integration (SAP, QuickBooks)",
        "Bank statement and transaction digitization",
    ], accent=ROSE)

    # ═══════════════════════════════════════════════════════
    # SLIDE 11: SELF-HEALING — THE KEY DIFFERENTIATOR
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Key Innovation", "🔄 Self-Healing Error Resolution",
                   "The Error Resolver Agent automatically detects and fixes mistakes — no human review needed.")
    footer(s)

    # Before/After comparison
    card(s, Inches(0.6), Inches(2.1), Inches(3.8), Inches(4.5), "❌", "BEFORE (Error Detected)", [
        "Invoice OCR reads: Grand Total = $2,750.00",
        "But Subtotal ($250) + Tax ($25) = $275.00",
        "Math audit FAILS: $275 ≠ $2,750",
        "",
        "License plate reads: ABO2CDE",
        "Jurisdiction regex requires digits at pos 2-3",
        "Syntax validation FAILS",
    ], accent=ROSE, bullet_size=Pt(11))

    # Arrow between
    arrow_right(s, Inches(4.65), Inches(4.0), Inches(0.8), EMERALD)
    tb = s.shapes.add_textbox(Inches(4.5), Inches(4.45), Inches(1.1), Inches(0.3))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Auto-Fix"; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = EMERALD; p.font.name = "Calibri"

    card(s, Inches(5.7), Inches(2.1), Inches(3.8), Inches(4.5), "✅", "AFTER (Self-Healed)", [
        "Detects decimal shift: $2,750 → $275.00",
        "Recalculates: $250 + $25 = $275.00  ✓",
        "Math audit PASSES",
        "",
        "Detects char confusion: 'O' should be '0'",
        "Fixes: ABO2CDE → AB02CDE",
        "Syntax validation PASSES",
    ], accent=EMERALD, bullet_size=Pt(11))

    # Third strategy box
    card(s, Inches(9.8), Inches(2.1), Inches(3.0), Inches(4.5), "🖼️", "Image Re-Filter Loop", [
        "If image is too dark or blurry:",
        "→ Requests CLAHE boost",
        "→ Inverts colors",
        "→ Loops back to Preprocessor",
        "→ Re-scans the image",
        "→ Retries up to max limit",
        "→ Always returns best result",
    ], accent=PURPLE, bullet_size=Pt(11))

    # ═══════════════════════════════════════════════════════
    # SLIDE 12: SELF-HEALING FLOW DIAGRAM
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Self-Healing Flow", "How the Error Resolver Makes Decisions",
                   "A step-by-step decision tree showing how the Error Resolver diagnoses and fixes each type of error.")
    footer(s)

    # Decision flow as cards
    # Start
    step_box(s, Inches(0.8), Inches(2.0), Inches(1.8), Inches(0.5), "!", "Error Detected", ROSE, "Validation failed")
    arrow_right(s, Inches(2.5), Inches(2.18), Inches(0.5), BORDER)

    # Decision: retries exceeded?
    dec = s.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(3.3), Inches(1.9), Inches(2.0), Inches(1.0))
    dec.fill.solid(); dec.fill.fore_color.rgb = RGBColor(254, 243, 199); dec.line.color.rgb = AMBER
    dtf = dec.text_frame; dtf.word_wrap = True; dtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    dtf.paragraphs[0].text = "Retries left?"; dtf.paragraphs[0].font.size = Pt(11)
    dtf.paragraphs[0].font.bold = True; dtf.paragraphs[0].font.name = "Calibri"

    # YES path → diagnose
    arrow_right(s, Inches(5.5), Inches(2.18), Inches(0.5), EMERALD)
    tb = s.shapes.add_textbox(Inches(5.55), Inches(1.85), Inches(0.5), Inches(0.3))
    tf = tb.text_frame; tf.paragraphs[0].text = "Yes"
    tf.paragraphs[0].font.size = Pt(9); tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = EMERALD

    # NO path → graceful exit
    arrow_down(s, Inches(4.2), Inches(3.0), Inches(0.4), ROSE)
    tb = s.shapes.add_textbox(Inches(4.5), Inches(3.0), Inches(0.5), Inches(0.3))
    tf = tb.text_frame; tf.paragraphs[0].text = "No"
    tf.paragraphs[0].font.size = Pt(9); tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ROSE
    grace = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.3), Inches(3.5), Inches(2.0), Inches(0.6))
    grace.fill.solid(); grace.fill.fore_color.rgb = RGBColor(254, 226, 226); grace.line.color.rgb = ROSE
    gtf = grace.text_frame; gtf.word_wrap = True; gtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    gtf.paragraphs[0].text = "Best-Effort Output"; gtf.paragraphs[0].font.size = Pt(10)
    gtf.paragraphs[0].font.bold = True; gtf.paragraphs[0].font.name = "Calibri"

    # Three strategy boxes
    strategies = [
        (Inches(6.3), "🧮", "Math Fix", "Rebalance equation\n$2750→$275.00\nMark valid ✓", AMBER),
        (Inches(8.8), "🔤", "Char Swap", "Fix O→0, I→1\nRe-validate syntax\nMark valid ✓", EMERALD),
        (Inches(11.0), "🖼️", "Re-Filter", "Apply CLAHE/Invert\nLoop to Preprocessor\nRe-scan image", PURPLE),
    ]
    for x, icon, title, desc, color in strategies:
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), Inches(1.9), Inches(2.5))
        shape.fill.solid(); shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = color; shape.line.width = Pt(2)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.9), Inches(1.9), Inches(0.07))
        bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
        tb = s.shapes.add_textbox(x + Inches(0.15), Inches(2.1), Inches(1.6), Inches(0.4))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = f"{icon} {title}"
        p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = NAVY; p.font.name = "Calibri"
        tb2 = s.shapes.add_textbox(x + Inches(0.15), Inches(2.6), Inches(1.6), Inches(1.5))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2 = tf2.paragraphs[0]; p2.text = desc
        p2.font.size = Pt(10); p2.font.color.rgb = TEXT_BODY; p2.font.name = "Calibri"
        p2.line_spacing = Pt(15)

    # ═══════════════════════════════════════════════════════
    # SLIDE 13: TECHNOLOGY STACK
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Technology", "Technology Stack & Architecture",
                   "Built entirely in Python with production-grade open-source libraries.")
    footer(s)

    tech_items = [
        ("🔗", "LangGraph", "StateGraph engine for multi-agent orchestration with conditional edges and cyclical loops.", MID_BLUE),
        ("👁️", "RapidOCR", "ONNX-based OCR engine for fast CPU/GPU text detection with bounding box coordinates.", CYAN),
        ("🖼️", "OpenCV", "Computer vision preprocessing: deskew (HoughLines), CLAHE contrast, bilateral noise filtering.", EMERALD),
        ("📊", "Streamlit", "Interactive web dashboard for file upload, image preview, and real-time agent monitoring.", ROSE),
        ("📄", "PyMuPDF", "PDF rendering at 200 DPI with digital text extraction for multi-format document loading.", AMBER),
        ("🧪", "Pytest", "Automated test suite with 9 passing tests covering all agents and self-healing scenarios.", PURPLE),
    ]
    for i, (icon, name, desc, color) in enumerate(tech_items):
        col = i % 3
        row = i // 3
        x = Inches(0.6) + Inches(4.1) * col
        y = Inches(2.1) + Inches(2.5) * row
        card(s, x, y, Inches(3.8), Inches(2.1), icon, name, [desc], accent=color)

    # ═══════════════════════════════════════════════════════
    # SLIDE 14: TESTING & VERIFICATION
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    section_header(s, "Quality Assurance", "Testing & Verification Results",
                   "Every agent and self-healing strategy is covered by automated tests.")
    footer(s)

    # Stats row
    stats = [
        ("9 / 9", "Tests Passing", EMERALD),
        ("100%", "Pass Rate", MID_BLUE),
        ("3", "Self-Healing\nStrategies Tested", AMBER),
        ("< 1s", "Avg Agent\nExecution Time", PURPLE),
    ]
    for i, (val, label, color) in enumerate(stats):
        x = Inches(0.8) + Inches(3.1) * i
        stat_block(s, x, Inches(2.0), val, label, color)

    # Test list
    tests = [
        ("✅", "test_orchestrator_classification", "Verifies correct routing for all 3 document types"),
        ("✅", "test_error_resolver_invoice_math_healing", "Validates automatic decimal shift correction"),
        ("✅", "test_error_resolver_plate_confusion_healing", "Tests O→0 and I→1 character disambiguation"),
        ("✅", "test_end_to_end_document_graph", "Full pipeline: image → preprocessor → OCR → digitizer → JSON"),
        ("✅", "test_end_to_end_invoice_graph_with_error_resolution", "Invoice with intentional math error → self-healed output"),
        ("✅", "test_end_to_end_license_plate", "Full pipeline: plate image → ANPR → validated plate number"),
    ]
    for i, (icon, name, desc) in enumerate(tests):
        y = Inches(3.5) + Inches(0.55) * i
        tb = s.shapes.add_textbox(Inches(0.8), y, Inches(0.4), Inches(0.35))
        tf = tb.text_frame; tf.paragraphs[0].text = icon; tf.paragraphs[0].font.size = Pt(14)
        tb2 = s.shapes.add_textbox(Inches(1.3), y, Inches(4.5), Inches(0.35))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p = tf2.paragraphs[0]; p.text = name
        p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = NAVY; p.font.name = "Consolas"
        tb3 = s.shapes.add_textbox(Inches(6.0), y, Inches(6.5), Inches(0.35))
        tf3 = tb3.text_frame; tf3.word_wrap = True
        p3 = tf3.paragraphs[0]; p3.text = desc
        p3.font.size = Pt(11); p3.font.color.rgb = TEXT_BODY; p3.font.name = "Calibri"

    # ═══════════════════════════════════════════════════════
    # SLIDE 15: THANK YOU / CONTACT
    # ═══════════════════════════════════════════════════════
    s = prs.slides.add_slide(blank)
    fill_bg(s, NAVY)

    tb = s.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "Thank You"
    p.font.size = Pt(48); p.font.bold = True; p.font.color.rgb = WHITE
    p.font.name = "Calibri"; p.alignment = PP_ALIGN.CENTER

    tb2 = s.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(11.3), Inches(0.6))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    p2 = tf2.paragraphs[0]; p2.text = "Multi-Agent OCR Platform with Self-Healing AI"
    p2.font.size = Pt(18); p2.font.color.rgb = SKY
    p2.font.name = "Calibri"; p2.alignment = PP_ALIGN.CENTER

    # Contact details
    contact = [
        ("👤", "Pankaj Todkar"),
        ("🔗", "github.com/pankajatd/ocr-multiagent-langgraph-platform"),
        ("📧", "pankajatd@gmail.com"),
    ]
    for i, (icon, text) in enumerate(contact):
        y = Inches(4.3) + Inches(0.55) * i
        tb = s.shapes.add_textbox(Inches(3.5), y, Inches(6.3), Inches(0.4))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = f"{icon}   {text}"; p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(15); p.font.color.rgb = TEXT_LIGHT; p.font.name = "Calibri"

    # ── Save ──
    out_path = os.path.join(os.getcwd(), output_filename)
    prs.save(out_path)
    print(f"\nPresentation created: {out_path}")
    print(f"Slides: {len(prs.slides)}")
    return out_path


if __name__ == "__main__":
    create_deck()
